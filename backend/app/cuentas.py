"""Parser de archivos de Cuentas fijas (Luz, Agua, Dividendo, etc.)."""
import datetime
import openpyxl


def parse_amount(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return 0 if val != val else float(val)
    s = str(val).strip()
    if s in ('', '-', 'N/A'):
        return 0
    s = s.replace('$', '').replace(' ', '')
    if '.' in s and ',' in s:
        s = s.replace('.', '').replace(',', '.')
    elif '.' in s:
        parts = s.split('.')
        if len(parts) == 2 and len(parts[1]) == 3:
            s = s.replace('.', '')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0


def parse_cuentas_xlsx(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    monthly_totals = {}
    details = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        monthly = {}
        payment_rows = []
        last_date = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 2:
                continue
            date_val, amount_val = row[0], row[1]
            if isinstance(date_val, datetime.datetime):
                last_date = date_val
            elif date_val is not None and str(date_val).strip() not in ('', '-'):
                try:
                    last_date = datetime.datetime.strptime(str(date_val).strip(), '%Y-%m-%d')
                except ValueError:
                    pass
            if last_date is None:
                continue
            amount = parse_amount(amount_val)
            if amount == 0:
                continue
            key = last_date.strftime('%Y-%m')
            monthly[key] = monthly.get(key, 0) + amount
            payment_rows.append({
                "date": last_date.strftime('%Y-%m-%d'),
                "month": key,
                "amount": amount,
            })
        monthly_totals[sheet_name] = monthly
        details[sheet_name] = payment_rows
    return monthly_totals, details
